"""Resource-bounded local document parsers. No OCR, macros or remote fetch."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile
import csv
import hashlib
import re


class DocumentError(ValueError):
    pass


class DocumentLimitError(DocumentError):
    pass


@dataclass(frozen=True)
class DocumentLimits:
    max_bytes: int = 20_000_000
    max_characters: int = 1_000_000
    max_pages: int = 500
    max_rows: int = 20_000
    max_columns: int = 100
    max_sheets: int = 30
    max_zip_members: int = 5000
    max_uncompressed_zip_bytes: int = 80_000_000


@dataclass(frozen=True)
class DocumentSection:
    text: str
    section: str
    page: int | None = None
    sheet: str | None = None


@dataclass(frozen=True)
class ParsedDocument:
    filename: str
    sha256: str
    mime_type: str
    sections: tuple[DocumentSection, ...]
    character_count: int


_MIMES = {
    ".txt": "text/plain", ".md": "text/markdown", ".markdown": "text/markdown",
    ".csv": "text/csv", ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_ENERGY_QUANTITY = re.compile(
    r"(?<![\w.])(?P<value>[+-]?\d+(?:[.,]\d+)?)\s*"
    r"(?P<unit>kWh|MWh|Wh|kW|MW|W|kVA|VA|kV|V|mA|A|Hz|°C|%)(?![A-Za-z])",
    re.IGNORECASE)
_MONEY = re.compile(r"(?P<symbol>₦|\$)\s*(?P<value>[\d,]+(?:\.\d+)?)")


def energy_facts(text: str, *, limit: int = 30) -> tuple[dict, ...]:
    """Only explicitly observed quantities; never interpret as certified settings."""
    facts: list[dict] = []
    for match in _ENERGY_QUANTITY.finditer(text):
        facts.append({"value": float(match.group("value").replace(",", "")),
                      "unit": match.group("unit"), "evidence": match.group()})
        if len(facts) >= limit:
            return tuple(facts)
    for match in _MONEY.finditer(text):
        facts.append({"value": float(match.group("value").replace(",", "")),
                      "unit": "NGN" if match.group("symbol") == "₦" else "USD",
                      "evidence": match.group()})
        if len(facts) >= limit:
            break
    return tuple(facts)


def _office_zip_guard(data: bytes, limits: DocumentLimits) -> None:
    try:
        with ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > limits.max_zip_members:
                raise DocumentLimitError("too many Office archive entries")
            total = 0
            for info in infos:
                total += info.file_size
                if total > limits.max_uncompressed_zip_bytes:
                    raise DocumentLimitError("Office decompression budget exceeded")
                if info.file_size > 100_000 and info.file_size > info.compress_size*1000:
                    raise DocumentLimitError("suspicious Office compression ratio")
                if info.flag_bits & 1 or info.filename.lower().endswith(("vbaproject.bin", ".vbs", ".exe")):
                    raise DocumentError("encrypted or macro-enabled Office input rejected")
    except BadZipFile:
        raise DocumentError("invalid Office archive") from None


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise DocumentError("only UTF-8 text is supported") from None


def _pdf(data: bytes, limits: DocumentLimits) -> list[DocumentSection]:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise DocumentError("PDF requires pip install 'energychain[documents]'") from None
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise DocumentError("encrypted PDF is not supported")
        if len(reader.pages) > limits.max_pages:
            raise DocumentLimitError("PDF page count exceeded")
        return [DocumentSection(txt, f"Page {i}", page=i)
                for i, page in enumerate(reader.pages, 1)
                if (txt := (page.extract_text() or "")).strip()]
    except DocumentError:
        raise
    except Exception:
        raise DocumentError("PDF extraction failed") from None


def _docx(data: bytes, limits: DocumentLimits) -> list[DocumentSection]:
    _office_zip_guard(data, limits)
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        raise DocumentError("DOCX requires pip install 'energychain[documents]'") from None
    try:
        document = Document(BytesIO(data))
        heading = "Document"
        output = []
        for part in document.iter_inner_content():
            if isinstance(part, Paragraph):
                text = part.text.strip()
                if not text:
                    continue
                if part.style and part.style.name.startswith("Heading"):
                    heading = text[:160]
                output.append(DocumentSection(text, heading))
            elif isinstance(part, Table):
                for row_id, row in enumerate(part.rows, 1):
                    if row_id > limits.max_rows:
                        raise DocumentLimitError("DOCX table exceeds row limit")
                    cells = [c.text.replace("\n", " ").strip()[:1000] for c in row.cells]
                    if any(cells):
                        output.append(DocumentSection(" | ".join(cells),
                                                      f"{heading} / table row {row_id}"))
        return output
    except DocumentError:
        raise
    except Exception:
        raise DocumentError("DOCX extraction failed") from None


def _xlsx(data: bytes, limits: DocumentLimits) -> list[DocumentSection]:
    _office_zip_guard(data, limits)
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise DocumentError("XLSX requires pip install 'energychain[documents]'") from None
    workbook = None
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
        if len(workbook.sheetnames) > limits.max_sheets:
            raise DocumentLimitError("spreadsheet sheet count exceeded")
        output = []
        for sheet in workbook:
            if (sheet.max_row and sheet.max_row > limits.max_rows or
                    sheet.max_column and sheet.max_column > limits.max_columns):
                raise DocumentLimitError("spreadsheet dimensions exceeded")
            for row_id, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_id > limits.max_rows or len(row) > limits.max_columns:
                    raise DocumentLimitError("spreadsheet row or column limit exceeded")
                values = [str(v).strip()[:1000] if v is not None else "" for v in row]
                if any(values):
                    output.append(DocumentSection(" | ".join(values),
                                                  f"row {row_id}", sheet=sheet.title))
        return output
    except DocumentError:
        raise
    except Exception:
        raise DocumentError("spreadsheet extraction failed") from None
    finally:
        if workbook is not None:
            workbook.close()


def parse_document(data: bytes, filename: str, *,
                   limits: DocumentLimits | None = None) -> ParsedDocument:
    limits = limits or DocumentLimits()
    if not isinstance(data, bytes) or not 0 < len(data) <= limits.max_bytes:
        raise DocumentLimitError("empty, invalid or oversized document")
    filename = Path(filename).name
    if not filename or len(filename) > 255:
        raise DocumentError("invalid filename")
    ext = Path(filename).suffix.lower()
    if ext not in _MIMES:
        raise DocumentError("unsupported document type")
    if ext == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise DocumentError("invalid PDF signature")
        sections = _pdf(data, limits)
    elif ext == ".docx":
        sections = _docx(data, limits)
    elif ext == ".xlsx":
        sections = _xlsx(data, limits)
    elif ext == ".csv":
        sections = []
        try:
            csv.field_size_limit(500_000)
            for i, row in enumerate(csv.reader(StringIO(_decode(data))), 1):
                if i > limits.max_rows or len(row) > limits.max_columns:
                    raise DocumentLimitError("CSV row or column limit exceeded")
                if any(c.strip() for c in row):
                    sections.append(DocumentSection(" | ".join(c.strip() for c in row),
                                                    f"CSV row {i}"))
        except csv.Error:
            raise DocumentError("invalid CSV data") from None
    else:
        text = _decode(data)
        if ext in (".md", ".markdown"):
            sections = []
            heading = "Document"
            for part in re.split(r"(?=^#{1,6}\s)", text, flags=re.M):
                if not part.strip():
                    continue
                first = part.split("\n", 1)[0]
                if re.match(r"^#{1,6}\s", first):
                    heading = first.lstrip("# ").strip()[:160]
                sections.append(DocumentSection(part.strip(), heading))
        else:
            sections = [DocumentSection(text, "Document")]
    chars = 0
    for sec in sections:
        chars += len(sec.text)
        if chars > limits.max_characters:
            raise DocumentLimitError("extracted text limit exceeded")
    if not chars:
        raise DocumentError("no extractable text (OCR not included)")
    return ParsedDocument(filename, hashlib.sha256(data).hexdigest(),
                          _MIMES[ext], tuple(sections), chars)


def chunk_document(doc: ParsedDocument, *, chunk_chars: int = 1200,
                   overlap_chars: int = 160) -> tuple[tuple[str, DocumentSection, int], ...]:
    if not 200 <= chunk_chars <= 10_000 or not 0 <= overlap_chars < chunk_chars//2:
        raise ValueError("invalid chunk sizing")
    chunks = []
    for section in doc.sections:
        txt = section.text.strip()
        for start in range(0, len(txt), chunk_chars-overlap_chars):
            fragment = txt[start:start+chunk_chars].strip()
            if fragment:
                chunks.append((fragment, section, start))
            if start+chunk_chars >= len(txt):
                break
    return tuple(chunks)
